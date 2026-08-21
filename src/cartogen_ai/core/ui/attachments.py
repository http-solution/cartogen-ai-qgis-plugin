# -*- coding: utf-8 -*-
"""
File-attachment reading for the chat dock's "attach a file" flow.

Extracted out of ui/dock_widget.py (docs/DESTRUCTIVE_TOOLS_AUDIT_2026-08-21.md's sibling
task, tracked as part of docs/ENGINEERING_PRODUCT_UX_REVIEW_2026-08-20.md SS3.2's
dock_widget.py-split recommendation) specifically because this function has zero Qt/QGIS
dependencies -- unlike the rest of dock_widget.py, which imports `qgis.PyQt`/`qgis.core`
unconditionally at module level with no QGIS_AVAILABLE-style fallback (see that file's
_extract_theme_palette docstring), and so cannot be imported or unit tested outside a real
QGIS process at all. Pulling this one function out into its own Qt-free module makes it
directly testable for the first time -- see tests/test_attachments.py -- following the same
"keep Qt-free logic separately testable" split CONTRIBUTING.md documents for
ui/chat_formatting.py.

This is a deliberately narrow, low-risk first step on the larger dock_widget.py-split
recommendation, not the whole of it. The much larger, higher-risk part -- splitting
CartogenAiDockWidget itself into separate ChatTabWidget/TasksTabWidget classes -- is
NOT done here. dock_widget.py cannot be imported in this sandbox at all (see above), so
there is no way to run or exercise a refactor of the live widget class beyond a syntax
check (py_compile) and manual cross-reference reading -- not a real verification for a
stateful Qt widget with cross-tab signal wiring, shared agent/task_manager references,
and background-thread callbacks. Attempting that split blind, with no way to catch a
broken signal connection or a missing attribute before it ships, is a worse outcome than
leaving it as one file a bit longer. See the "Splitting CartogenAiDockWidget itself"
section of docs/ENGINEERING_PRODUCT_UX_REVIEW_2026-08-20.md's own SS3.2 finding for the
scoped plan this deferral is tracked against -- it should be done in (or immediately
verified in) a real QGIS session, not this sandbox.
"""

import os


def read_attached_file(path):
    """Reads a file attached via the dock's file-picker and returns
    (data, error) where data is a dict describing the content for the model
    ({'text': ..., 'is_image': bool} or {'text': None, 'is_image': True, 'b64':
    ..., 'mime': ...} for images) and error is a user-facing string on
    failure, or None on success -- exactly one of the two is non-None. Original
    implementation, byte-identical logic, just relocated -- see this module's
    docstring for why."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return {"text": text, "is_image": False}, None

        if ext == ".docx":
            from docx import Document
            doc = Document(path)
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\t"
                    text += "\n"
            return {"text": text, "is_image": False}, None

        if ext in (".png", ".jpg", ".jpeg", ".bmp", ".gif"):
            import base64
            with open(path, "rb") as f:
                b64 = base64.b64encode(f.read()).decode()
            mime = "jpeg" if ext in (".jpg", ".jpeg") else ext.lstrip(".")
            return {"text": None, "is_image": True, "b64": b64, "mime": mime}, None

        if ext == ".csv":
            import pandas as pd
            df = None
            last_err = None
            # Excel-exported CSVs are often cp1252 (curly quotes, em dashes),
            # not UTF-8 -- try in order of specificity and fall back to
            # latin-1 last since it never raises a decode error itself.
            for enc in ("utf-8-sig", "cp1252", "latin-1"):
                try:
                    df = pd.read_csv(path, encoding=enc)
                    break
                except UnicodeDecodeError as e:
                    last_err = e
            if df is None:
                raise last_err
            text = (
                f"CSV with {len(df)} rows, {len(df.columns)} columns\n"
                f"Columns: {list(df.columns)}\n"
                f"First 5 rows:\n{df.head().to_string()}"
            )
            return {"text": text, "is_image": False}, None

        if ext in (".xlsx", ".xls"):
            import pandas as pd
            df = pd.read_excel(path)
            text = (
                f"Excel with {len(df)} rows, {len(df.columns)} columns\n"
                f"Columns: {list(df.columns)}\n"
                f"First 5 rows:\n{df.head().to_string()}"
            )
            return {"text": text, "is_image": False}, None

        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return {"text": f.read(), "is_image": False}, None
    except ImportError as e:
        return None, (
            f"Required package missing: {e.name}. Run in OSGeo4W Shell: "
            f"python -m pip install pypdf python-docx openpyxl pandas"
        )
    except Exception as e:
        return None, f"Error: {e}"
