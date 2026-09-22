"""Build deterministic, disposable inputs for docs/RELEASE_LIVE_TEST_SCENARIOS.md.

Run --documents with ordinary Python, then --spatial with QGIS's python-qgis.bat.
The generated files are intentionally small and contain synthetic data only.
"""

import argparse
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1] / "docs" / "release_smoke_assets"
INPUTS = ROOT / "inputs"


def ensure_folders():
    for name in ("inputs", "outputs", "screenshots", "logs"):
        (ROOT / name).mkdir(parents=True, exist_ok=True)


def build_documents():
    import pymupdf
    from docx import Document

    ensure_folders()
    rows = [("id", "value"), ("A01", "10"), ("A02", "20"), ("A03", "30")]

    doc = Document()
    doc.add_heading("Synthetic smoke-test table", 0)
    table = doc.add_table(rows=len(rows), cols=2)
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value
    doc.save(INPUTS / "smoke_tables.docx")

    pdf = pymupdf.open()
    page = pdf.new_page(width=350, height=260)
    page.insert_text((35, 35), "Synthetic smoke-test table", fontsize=13)
    left, top, width, row_height = 35, 55, 240, 34
    for index in range(len(rows) + 1):
        y = top + index * row_height
        page.draw_line((left, y), (left + width, y))
    for x in (left, left + width / 2, left + width):
        page.draw_line((x, top), (x, top + len(rows) * row_height))
    for index, (item_id, value) in enumerate(rows):
        y = top + index * row_height + 22
        page.insert_text((left + 8, y), item_id, fontsize=11)
        page.insert_text((left + width / 2 + 8, y), value, fontsize=11)
    pdf.save(INPUTS / "smoke_tables.pdf")
    pdf.close()
    print("Created smoke_tables.pdf and smoke_tables.docx")


def build_spatial():
    from osgeo import gdal, ogr, osr
    from qgis.core import QgsApplication, QgsCoordinateReferenceSystem, QgsProject, QgsRasterLayer, QgsVectorLayer

    ensure_folders()
    gdal.UseExceptions()
    gpkg_path = INPUTS / "smoke_data.gpkg"
    if gpkg_path.exists():
        gpkg_path.unlink()
    driver = ogr.GetDriverByName("GPKG")
    data = driver.CreateDataSource(str(gpkg_path))
    srs = osr.SpatialReference()
    srs.ImportFromEPSG(32636)
    wgs84 = osr.SpatialReference()
    wgs84.ImportFromEPSG(4326)
    wgs84.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    srs.SetAxisMappingStrategy(osr.OAMS_TRADITIONAL_GIS_ORDER)
    transform = osr.CoordinateTransformation(wgs84, srs)

    def xy(lon, lat):
        east, north, _ = transform.TransformPoint(lon, lat)
        return east, north

    def layer(name, geom_type, fields):
        created = data.CreateLayer(name, srs, geom_type=geom_type)
        for field_name, kind in fields:
            created.CreateField(ogr.FieldDefn(field_name, kind))
        return created

    def add_point(target, lon, lat, values):
        feature = ogr.Feature(target.GetLayerDefn())
        point = ogr.Geometry(ogr.wkbPoint)
        east, north = xy(lon, lat)
        point.AddPoint_2D(east, north)
        feature.SetGeometry(point)
        for key, value in values.items():
            feature.SetField(key, value)
        target.CreateFeature(feature)

    def add_rect(target, west, south, east, north, values):
        ring = ogr.Geometry(ogr.wkbLinearRing)
        for lon, lat in ((west, south), (east, south), (east, north), (west, north), (west, south)):
            ring.AddPoint_2D(*xy(lon, lat))
        polygon = ogr.Geometry(ogr.wkbPolygon)
        polygon.AddGeometry(ring)
        feature = ogr.Feature(target.GetLayerDefn())
        feature.SetGeometry(polygon)
        for key, value in values.items():
            feature.SetField(key, value)
        target.CreateFeature(feature)

    points = layer("smoke_points", ogr.wkbPoint, [("name", ogr.OFTString), ("category", ogr.OFTString), ("population", ogr.OFTInteger), ("need", ogr.OFTInteger)])
    point_rows = [
        (35.922, 31.948, "North", "clinic", 1200, 3),
        (35.929, 31.949, "Central", "school", 2100, 8),
        (35.936, 31.953, "East", "clinic", 1650, 6),
        (35.925, 31.958, "West", "shelter", 950, 9),
        (35.942, 31.961, "South", "school", 1800, 4),
    ]
    for lon, lat, name, category, population, need in point_rows:
        add_point(points, lon, lat, dict(name=name, category=category, population=population, need=need))

    hubs = layer("smoke_hubs", ogr.wkbPoint, [("name", ogr.OFTString)])
    for lon, lat, name in ((35.927, 31.950, "Hub_A"), (35.936, 31.958, "Hub_B"), (35.918, 31.961, "Hub_C")):
        add_point(hubs, lon, lat, dict(name=name))

    boundary = layer("smoke_boundary", ogr.wkbPolygon, [("admin_name", ogr.OFTString), ("admin_type", ogr.OFTString)])
    add_rect(boundary, 35.915, 31.941, 35.950, 31.969, dict(admin_name="Smoke Area", admin_type="test boundary"))

    zones = layer("smoke_zones", ogr.wkbPolygon, [("zone_type", ogr.OFTString)])
    add_rect(zones, 35.916, 31.942, 35.927, 31.968, dict(zone_type="urban"))
    add_rect(zones, 35.927, 31.942, 35.938, 31.968, dict(zone_type="mixed"))
    add_rect(zones, 35.938, 31.942, 35.949, 31.968, dict(zone_type="open"))

    admin = layer("smoke_admin", ogr.wkbPolygon, [("admin_name", ogr.OFTString), ("population", ogr.OFTInteger), ("need", ogr.OFTInteger)])
    for index, (west, east, population, need) in enumerate(((35.916, 35.927, 1000, 3), (35.927, 35.938, 2000, 7), (35.938, 35.949, 1500, 10)), start=1):
        add_rect(admin, west, 31.942, east, 31.968, dict(admin_name=f"District_{index}", population=population, need=need))
    data = None

    west, south = xy(35.915, 31.941)
    east, north = xy(35.950, 31.969)
    width = height = 128
    for name, bands in (("smoke_dem.tif", 1), ("smoke_image.tif", 3)):
        ds = gdal.GetDriverByName("GTiff").Create(str(INPUTS / name), width, height, bands, gdal.GDT_Byte)
        ds.SetGeoTransform((west, (east - west) / width, 0, north, 0, -(north - south) / height))
        ds.SetProjection(srs.ExportToWkt())
        import numpy as np
        yy, xx = np.mgrid[0:height, 0:width]
        for band in range(1, bands + 1):
            values = (40 + 0.5 * xx + 0.8 * yy + 28 * np.sin(xx / 13 + band)) % 255
            ds.GetRasterBand(band).WriteArray(values.astype("uint8"))
        ds = None

    app = QgsApplication([], False)
    app.initQgis()
    project = QgsProject.instance()
    project.setCrs(QgsCoordinateReferenceSystem("EPSG:32636"))
    for name in ("smoke_boundary", "smoke_zones", "smoke_admin", "smoke_hubs", "smoke_points"):
        vector = QgsVectorLayer(f"{gpkg_path}|layername={name}", name, "ogr")
        if not vector.isValid():
            raise RuntimeError(f"Cannot load {name}")
        project.addMapLayer(vector)
    for name in ("smoke_dem", "smoke_image"):
        raster = QgsRasterLayer(str(INPUTS / f"{name}.tif"), name)
        if not raster.isValid():
            raise RuntimeError(f"Cannot load {name}")
        project.addMapLayer(raster)
    project.setFileName(str(INPUTS / "smoke_start.qgz"))
    if not project.write():
        raise RuntimeError("Could not save smoke_start.qgz")
    project.clear()
    app.exitQgis()
    print("Created smoke_data.gpkg, two GeoTIFFs, and smoke_start.qgz")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--documents", action="store_true")
    parser.add_argument("--spatial", action="store_true")
    args = parser.parse_args()
    if args.documents:
        build_documents()
    elif args.spatial:
        build_spatial()
    else:
        parser.error("choose --documents or --spatial")
