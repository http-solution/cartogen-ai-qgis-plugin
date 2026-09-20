# -*- coding: utf-8 -*-
"""Tests for agent/tools/reporting_tools.py (generate_chart, extract_pdf_tables,
extract_word_tables, aggregate_data) -- the Phase 1 humanitarian-ops
improvements: statistical charting and structured document-table extraction,
neither of which existed anywhere in the codebase before this pass."""
import os
import unittest
from cartogen_ai.core.agent.tools.reporting_tools import (
    generate_chart, extract_pdf_tables, extract_word_tables, aggregate_data,
    _aggregate_rows, _coerce_number, load_3w_data, _read_tabular_rows,
    _aggregate_3w_presence, _is_missing, generate_sector_coverage_report,
    _apply_pie_slice_cap,
)


def _skip_if_missing(test_case, module_name):
    try:
        __import__(module_name)
    except ImportError:
        test_case.skipTest(f"{module_name} not installed")


class TestCoerceNumber(unittest.TestCase):
    def test_parses_numeric_strings(self):
        self.assertEqual(_coerce_number("42"), 42.0)
        self.assertEqual(_coerce_number("3.5"), 3.5)

    def test_returns_none_for_non_numeric(self):
        self.assertIsNone(_coerce_number("N/A"))
        self.assertIsNone(_coerce_number(None))
        self.assertIsNone(_coerce_number(""))


class TestAggregateRows(unittest.TestCase):
    """Pure Python, no QGIS or optional deps needed."""

    def test_count_groups(self):
        rows = [{"district": "Aden"}, {"district": "Aden"}, {"district": "Taiz"}]
        groups, skipped = _aggregate_rows(rows, "district", None, "count")
        self.assertEqual(groups, {"Aden": 2, "Taiz": 1})
        self.assertEqual(skipped, 0)

    def test_sum_groups(self):
        rows = [
            {"district": "Aden", "idps": "100"},
            {"district": "Aden", "idps": "50"},
            {"district": "Taiz", "idps": "30"},
        ]
        groups, skipped = _aggregate_rows(rows, "district", "idps", "sum")
        self.assertEqual(groups, {"Aden": 150.0, "Taiz": 30.0})

    def test_mean_groups(self):
        rows = [{"d": "A", "v": "10"}, {"d": "A", "v": "20"}]
        groups, _ = _aggregate_rows(rows, "d", "v", "mean")
        self.assertEqual(groups["A"], 15.0)

    def test_skips_rows_missing_group_field(self):
        rows = [{"district": "Aden", "idps": "10"}, {"idps": "20"}]
        groups, skipped = _aggregate_rows(rows, "district", "idps", "sum")
        self.assertEqual(groups, {"Aden": 10.0})
        self.assertEqual(skipped, 1)

    def test_skips_rows_with_non_numeric_value(self):
        rows = [{"district": "Aden", "idps": "N/A"}, {"district": "Aden", "idps": "10"}]
        groups, skipped = _aggregate_rows(rows, "district", "idps", "sum")
        self.assertEqual(groups, {"Aden": 10.0})
        self.assertEqual(skipped, 1)


class TestAggregateDataValidation(unittest.TestCase):
    def test_rejects_empty_rows(self):
        res = aggregate_data([], "district")
        self.assertIn("error", res)

    def test_rejects_unknown_agg(self):
        res = aggregate_data([{"d": "A"}], "d", agg="median")
        self.assertIn("error", res)
        self.assertIn("agg", res["error"])

    def test_requires_value_field_unless_count(self):
        res = aggregate_data([{"d": "A"}], "d", agg="sum")
        self.assertIn("error", res)
        self.assertIn("value_field", res["error"])

    def test_count_does_not_require_value_field(self):
        res = aggregate_data([{"d": "A"}, {"d": "B"}], "d", agg="count")
        self.assertTrue(res["success"])
        self.assertEqual(res["groups"], {"A": 1, "B": 1})

    def test_reports_no_usable_rows(self):
        res = aggregate_data([{"other": "x"}], "district", agg="count")
        self.assertIn("error", res)


class TestGenerateChartValidation(unittest.TestCase):
    def test_rejects_unknown_chart_type(self):
        res = generate_chart("scatter", "t", ["a"], [1])
        self.assertIn("error", res)
        self.assertIn("chart_type", res["error"])

    def test_rejects_mismatched_lengths(self):
        res = generate_chart("bar", "t", ["a", "b"], [1])
        self.assertIn("error", res)

    def test_rejects_empty_data(self):
        res = generate_chart("bar", "t", [], [])
        self.assertIn("error", res)

    def test_rejects_negative_values_for_pie(self):
        res = generate_chart("pie", "t", ["a", "b"], [10, -5])
        self.assertIn("error", res)
        self.assertIn("negative", res["error"])


class TestApplyPieSliceCap(unittest.TestCase):
    """Phase 4 (2026-09-19): a pie chart with too many slices becomes unreadable --
    pure Python, no matplotlib needed."""

    def test_within_cap_returned_unchanged(self):
        labels, values, note = _apply_pie_slice_cap(["a", "b", "c"], [1, 2, 3], max_slices=8)
        self.assertEqual(labels, ["a", "b", "c"])
        self.assertEqual(values, [1, 2, 3])
        self.assertIsNone(note)

    def test_over_cap_combines_smallest_into_other(self):
        labels = [f"cat{i}" for i in range(10)]
        values = list(range(10))  # cat0=0 ... cat9=9
        new_labels, new_values, note = _apply_pie_slice_cap(labels, values, max_slices=5)
        self.assertEqual(len(new_labels), 5)
        self.assertIn("Other", new_labels)
        self.assertIsNotNone(note)
        # The 6 smallest (cat0..cat5, values 0-5) collapse into Other = 15;
        # the 4 largest (cat6..cat9, values 6-9) are kept individually.
        other_index = new_labels.index("Other")
        self.assertEqual(new_values[other_index], sum(range(6)))
        self.assertEqual(set(new_labels) - {"Other"}, {"cat6", "cat7", "cat8", "cat9"})

    def test_exactly_at_cap_is_unchanged(self):
        labels = [f"c{i}" for i in range(8)]
        values = list(range(8))
        new_labels, new_values, note = _apply_pie_slice_cap(labels, values, max_slices=8)
        self.assertEqual(new_labels, labels)
        self.assertIsNone(note)


class TestGenerateChartRendersRealFile(unittest.TestCase):
    def setUp(self):
        _skip_if_missing(self, "matplotlib")

    def test_bar_chart_produces_a_real_png(self):
        res = generate_chart("bar", "Incidents by District", ["Aden", "Taiz"], [12, 7])
        try:
            self.assertTrue(res["success"])
            self.assertTrue(os.path.exists(res["output_path"]))
            self.assertGreater(os.path.getsize(res["output_path"]), 0)
        finally:
            if res.get("output_path") and os.path.exists(res["output_path"]):
                os.remove(res["output_path"])

    def test_pie_chart_produces_a_real_png(self):
        res = generate_chart("pie", "Funding by Cluster", ["Health", "WASH"], [40, 60])
        try:
            self.assertTrue(res["success"])
            self.assertTrue(os.path.exists(res["output_path"]))
        finally:
            if res.get("output_path") and os.path.exists(res["output_path"]):
                os.remove(res["output_path"])

    def test_default_dpi_is_300_not_150(self):
        res = generate_chart("bar", "t", ["a", "b"], [1, 2])
        try:
            self.assertTrue(res["success"])
            from PIL import Image
            with Image.open(res["output_path"]) as img:
                width_px, _ = img.size
            # figsize=(8,5) inches * dpi=300 = 2400px wide, vs the old dpi=150 -> 1200px --
            # a real, measurable difference in the actual rendered file, not just the
            # function argument being accepted.
            self.assertGreater(width_px, 2000)
        finally:
            if res.get("output_path") and os.path.exists(res["output_path"]):
                os.remove(res["output_path"])

    def test_custom_dpi_is_respected(self):
        res = generate_chart("bar", "t", ["a", "b"], [1, 2], dpi=100)
        try:
            self.assertTrue(res["success"])
            from PIL import Image
            with Image.open(res["output_path"]) as img:
                width_px, _ = img.size
            self.assertLess(width_px, 1000)
        finally:
            if res.get("output_path") and os.path.exists(res["output_path"]):
                os.remove(res["output_path"])

    def test_pie_chart_over_slice_cap_gets_a_warning_and_fewer_slices(self):
        labels = [f"cat{i}" for i in range(12)]
        values = list(range(1, 13))
        res = generate_chart("pie", "Many Categories", labels, values)
        try:
            self.assertTrue(res["success"], res)
            self.assertIn("warning", res)
            self.assertIn("Other", res["warning"])
        finally:
            if res.get("output_path") and os.path.exists(res["output_path"]):
                os.remove(res["output_path"])

    def test_custom_color_palette_is_accepted(self):
        res = generate_chart(
            "bar", "t", ["a", "b"], [1, 2], color_palette=["#123456", "#abcdef"],
        )
        try:
            self.assertTrue(res["success"], res)
            self.assertTrue(os.path.exists(res["output_path"]))
        finally:
            if res.get("output_path") and os.path.exists(res["output_path"]):
                os.remove(res["output_path"])


class TestExtractPdfTablesValidation(unittest.TestCase):
    def test_rejects_missing_file(self):
        res = extract_pdf_tables("/definitely/not/a/real/file.pdf")
        self.assertIn("error", res)
        self.assertIn("not found", res["error"])


class TestExtractPdfTablesRealFile(unittest.TestCase):
    def setUp(self):
        _skip_if_missing(self, "pdfplumber")
        _skip_if_missing(self, "matplotlib")

    def test_extracts_a_real_table_from_a_generated_pdf(self):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages

        path = "scratch_test_sitrep.pdf"
        fig, ax = plt.subplots(figsize=(6, 3))
        ax.axis("off")
        data = [["District", "IDPs"], ["Aden", "1200"], ["Taiz", "800"]]
        ax.table(cellText=data, loc="center", cellLoc="center")
        try:
            with PdfPages(path) as pdf:
                pdf.savefig(fig)
            plt.close(fig)

            res = extract_pdf_tables(path)
            self.assertTrue(res["success"])
            self.assertEqual(res["table_count"], 1)
            self.assertEqual(res["tables"][0]["columns"], ["District", "IDPs"])
            self.assertEqual(res["tables"][0]["row_count"], 2)
        finally:
            if os.path.exists(path):
                os.remove(path)


class TestExtractWordTablesValidation(unittest.TestCase):
    def test_rejects_missing_file(self):
        res = extract_word_tables("/definitely/not/a/real/file.docx")
        self.assertIn("error", res)
        self.assertIn("not found", res["error"])


class TestExtractWordTablesRealFile(unittest.TestCase):
    def setUp(self):
        _skip_if_missing(self, "docx")

    def test_extracts_a_real_table_from_a_generated_docx(self):
        from docx import Document

        path = "scratch_test_needs.docx"
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "District"
        table.rows[0].cells[1].text = "Households"
        row = table.add_row().cells
        row[0].text, row[1].text = "Aden", "300"
        try:
            doc.save(path)
            res = extract_word_tables(path)
            self.assertTrue(res["success"])
            self.assertEqual(res["table_count"], 1)
            self.assertEqual(res["tables"][0]["columns"], ["District", "Households"])
            self.assertEqual(res["tables"][0]["rows"], [{"District": "Aden", "Households": "300"}])
        finally:
            if os.path.exists(path):
                os.remove(path)


class TestIsMissing(unittest.TestCase):
    def test_none_and_empty_string(self):
        self.assertTrue(_is_missing(None))
        self.assertTrue(_is_missing(""))

    def test_nan(self):
        self.assertTrue(_is_missing(float("nan")))

    def test_real_values_not_missing(self):
        self.assertFalse(_is_missing("Aden"))
        self.assertFalse(_is_missing(0))
        self.assertFalse(_is_missing("0"))


class TestReadTabularRowsCsv(unittest.TestCase):
    def test_reads_csv_rows_and_columns(self):
        path = "scratch_test_3w_rows.csv"
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write("org,district,sector\nNRC,Aden,Shelter\nUNICEF,Aden,WASH\nNRC,Taiz,Shelter\n")
        try:
            rows, columns = _read_tabular_rows(path)
            self.assertEqual(columns, ["org", "district", "sector"])
            self.assertEqual(len(rows), 3)
            self.assertEqual(rows[0], {"org": "NRC", "district": "Aden", "sector": "Shelter"})
        finally:
            os.remove(path)

    def test_rejects_unsupported_extension(self):
        with self.assertRaises(ValueError):
            _read_tabular_rows("file.txt")


class TestAggregate3wPresence(unittest.TestCase):
    """Pure Python, no QGIS needed -- mirrors the humanitarian 3W/4W ('who does
    what where') schema: one row per organization's activity in a district."""

    def test_counts_distinct_orgs_not_activity_rows(self):
        rows = [
            {"org": "NRC", "district": "Aden", "sector": "Shelter"},
            {"org": "NRC", "district": "Aden", "sector": "WASH"},  # same org, 2nd activity
            {"org": "UNICEF", "district": "Aden", "sector": "WASH"},
            {"org": "IRC", "district": "Taiz", "sector": "Health"},
        ]
        units, skipped = _aggregate_3w_presence(rows, "district", "org", "sector")
        self.assertEqual(skipped, 0)
        aden = units["aden"]
        self.assertEqual(aden["unit"], "Aden")
        self.assertEqual(aden["orgs"], {"NRC", "UNICEF"})  # 2 distinct orgs, not 3 rows
        self.assertEqual(aden["sectors"], {"Shelter", "WASH"})
        self.assertEqual(aden["activity_count"], 3)
        self.assertEqual(units["taiz"]["orgs"], {"IRC"})

    def test_normalizes_key_case_and_whitespace_but_keeps_display_casing(self):
        rows = [{"org": "NRC", "district": " Aden "}, {"org": "UNICEF", "district": "aden"}]
        units, _ = _aggregate_3w_presence(rows, "district", "org")
        self.assertEqual(len(units), 1)
        entry = units["aden"]
        self.assertEqual(entry["unit"], "Aden")  # first-seen casing kept
        self.assertEqual(entry["orgs"], {"NRC", "UNICEF"})

    def test_skips_rows_missing_admin_or_org(self):
        rows = [{"org": "NRC", "district": "Aden"}, {"org": "", "district": "Aden"}, {"org": "NRC", "district": ""}]
        units, skipped = _aggregate_3w_presence(rows, "district", "org")
        self.assertEqual(skipped, 2)
        self.assertEqual(len(units), 1)

    def test_nan_admin_or_org_skipped(self):
        rows = [{"org": float("nan"), "district": "Aden"}, {"org": "NRC", "district": "Aden"}]
        units, skipped = _aggregate_3w_presence(rows, "district", "org")
        self.assertEqual(skipped, 1)
        self.assertEqual(units["aden"]["orgs"], {"NRC"})

    def test_no_sector_field_omits_sectors(self):
        rows = [{"org": "NRC", "district": "Aden"}]
        units, _ = _aggregate_3w_presence(rows, "district", "org")
        self.assertEqual(units["aden"]["sectors"], set())


class TestLoad3wDataValidation(unittest.TestCase):
    def test_rejects_missing_file(self):
        res = load_3w_data("/definitely/not/a/real/file.csv", "district", "org")
        self.assertIn("error", res)
        self.assertIn("not found", res["error"])


class TestLoad3wDataRealFile(unittest.TestCase):
    def test_loads_and_aggregates_a_real_csv(self):
        path = "scratch_test_3w_load.csv"
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(
                "org,district,sector\n"
                "NRC,Aden,Shelter\n"
                "NRC,Aden,WASH\n"
                "UNICEF,Aden,WASH\n"
                "IRC,Taiz,Health\n"
            )
        try:
            res = load_3w_data(path, admin_field="district", org_field="org", sector_field="sector")
            self.assertTrue(res["success"])
            self.assertEqual(res["total_rows"], 4)
            self.assertEqual(res["distinct_admin_units"], 2)
            aden = next(u for u in res["units"] if u["unit"] == "Aden")
            self.assertEqual(aden["organization_count"], 2)
            self.assertEqual(aden["organizations"], ["NRC", "UNICEF"])
            self.assertEqual(aden["sector_count"], 2)
            self.assertEqual(aden["activity_count"], 3)
        finally:
            os.remove(path)

    def test_rejects_unknown_field_name(self):
        path = "scratch_test_3w_badfield.csv"
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write("org,district\nNRC,Aden\n")
        try:
            res = load_3w_data(path, admin_field="not_a_column", org_field="org")
            self.assertIn("error", res)
            self.assertIn("not found", res["error"])
        finally:
            os.remove(path)


class TestGenerateSectorCoverageReport(unittest.TestCase):
    """Thin composite over aggregate_data + generate_chart -- the humanitarian
    'beneficiaries reached vs. target, by sector' report table + bar chart in
    one call. Real charts are produced (matplotlib required), so cleaned up
    after each test."""

    def setUp(self):
        _skip_if_missing(self, "matplotlib")

    def test_computes_coverage_percent_and_sorts_worst_first(self):
        rows = [
            {"sector": "WASH", "reached": "500", "target": "1000"},
            {"sector": "WASH", "reached": "300", "target": "1000"},  # same sector, 2nd activity
            {"sector": "Health", "reached": "900", "target": "1000"},
            {"sector": "Shelter", "reached": "200", "target": "500"},
        ]
        res = generate_sector_coverage_report(rows, "sector", "reached", target_field="target")
        try:
            self.assertTrue(res.get("success"), res)
            table = res["table"]
            by_sector = {row["sector"]: row for row in table}
            self.assertEqual(by_sector["WASH"]["reached"], 800.0)
            self.assertEqual(by_sector["WASH"]["target"], 2000.0)
            self.assertEqual(by_sector["WASH"]["coverage_percent"], 40.0)
            self.assertEqual(by_sector["Health"]["coverage_percent"], 90.0)
            self.assertEqual(by_sector["Shelter"]["coverage_percent"], 40.0)
            # Worst (lowest) coverage first; Health (90%) last.
            self.assertEqual(table[-1]["sector"], "Health")
            self.assertLessEqual(table[0]["coverage_percent"], table[1]["coverage_percent"])
        finally:
            if res.get("chart_path") and os.path.exists(res["chart_path"]):
                os.remove(res["chart_path"])

    def test_missing_target_sorts_first_not_last(self):
        # A group with an unusable target value is a data gap worth flagging,
        # not something to bury at the bottom of a worst-coverage-first table.
        rows = [
            {"sector": "Health", "reached": "900", "target": "1000"},  # 90%
            {"sector": "Protection", "reached": "150", "target": "N/A"},  # target unusable
        ]
        res = generate_sector_coverage_report(rows, "sector", "reached", target_field="target")
        try:
            self.assertTrue(res.get("success"), res)
            table = res["table"]
            self.assertEqual(table[0]["sector"], "Protection")
            self.assertIsNone(table[0]["coverage_percent"])
            self.assertEqual(res["target_skipped_rows"], 1)
            self.assertEqual(res["reached_skipped_rows"], 0)
        finally:
            if res.get("chart_path") and os.path.exists(res["chart_path"]):
                os.remove(res["chart_path"])

    def test_without_target_field_sorts_by_reached_descending(self):
        rows = [
            {"sector": "WASH", "reached": "100"},
            {"sector": "Health", "reached": "500"},
            {"sector": "Shelter", "reached": "300"},
        ]
        res = generate_sector_coverage_report(rows, "sector", "reached")
        try:
            self.assertTrue(res.get("success"), res)
            table = res["table"]
            self.assertEqual([row["sector"] for row in table], ["Health", "Shelter", "WASH"])
            self.assertNotIn("target", table[0])
            self.assertNotIn("coverage_percent", table[0])
        finally:
            if res.get("chart_path") and os.path.exists(res["chart_path"]):
                os.remove(res["chart_path"])

    def test_propagates_aggregate_data_error(self):
        res = generate_sector_coverage_report([], "sector", "reached")
        self.assertIn("error", res)

    def test_chart_file_is_actually_created(self):
        rows = [{"sector": "WASH", "reached": "100"}, {"sector": "Health", "reached": "200"}]
        res = generate_sector_coverage_report(rows, "sector", "reached")
        try:
            self.assertTrue(res.get("success"), res)
            self.assertTrue(os.path.exists(res["chart_path"]))
        finally:
            if res.get("chart_path") and os.path.exists(res["chart_path"]):
                os.remove(res["chart_path"])


class TestExtractThenAggregatePipeline(unittest.TestCase):
    """Integration-style test confirming the two new tools chain together
    correctly, matching the intended workflow (extract a sitrep table, then
    aggregate it) rather than just testing each in isolation."""

    def setUp(self):
        _skip_if_missing(self, "docx")

    def test_word_table_feeds_directly_into_aggregate_data(self):
        from docx import Document

        path = "scratch_test_pipeline.docx"
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Priority"
        table.rows[0].cells[1].text = "Households"
        for pri, hh in [("High", "300"), ("High", "500"), ("Medium", "150")]:
            row = table.add_row().cells
            row[0].text, row[1].text = pri, hh
        try:
            doc.save(path)
            extracted = extract_word_tables(path)
            agg = aggregate_data(extracted["tables"][0]["rows"], "Priority", "Households", "sum")
            self.assertTrue(agg["success"])
            self.assertEqual(agg["groups"], {"High": 800.0, "Medium": 150.0})
        finally:
            if os.path.exists(path):
                os.remove(path)


if __name__ == "__main__":
    unittest.main()
