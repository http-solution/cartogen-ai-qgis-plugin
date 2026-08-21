# -*- coding: utf-8 -*-
"""Tests for ui/attachments.py's read_attached_file(). Extracted out of
ui/dock_widget.py in v1.2.31 specifically because it has zero Qt/QGIS
dependencies, unlike the rest of dock_widget.py which cannot be imported
outside a real QGIS process at all -- see ui/attachments.py's own docstring.
This is the first test coverage this function has ever had.

Uses tempfile (not a scratch_test_* file in the repo root) deliberately --
tests/test_reporting_tools.py's scratch_test_*.csv/.docx/.pdf fixtures hit a
known, pre-existing sandbox permission error on os.remove() cleanup (see
CHANGELOG.md's recurring "6 scratch-file-permission errors" baseline note).
tempfile.mkdtemp() writes outside the repo tree, where cleanup actually
works in this sandbox."""
import base64
import os
import shutil
import tempfile
import unittest

from cartogen_ai.core.ui.attachments import read_attached_file


def _skip_if_missing(test_case, module_name):
    try:
        __import__(module_name)
    except ImportError:
        test_case.skipTest(f"{module_name} not installed")


class TestReadAttachedFilePlainText(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_reads_a_plain_text_file(self):
        path = os.path.join(self.tmpdir, "notes.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("hello from a plain text attachment")
        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertEqual(data["text"], "hello from a plain text attachment")
        self.assertFalse(data["is_image"])

    def test_unreadable_bytes_fall_back_instead_of_raising(self):
        # errors="ignore" on the plain-text fallback path -- a file with
        # invalid UTF-8 bytes should degrade gracefully, not blow up the
        # whole attach-a-file flow.
        path = os.path.join(self.tmpdir, "binary.dat")
        with open(path, "wb") as f:
            f.write(b"\xff\xfe\x00\x01not valid utf-8")
        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertIsInstance(data["text"], str)

    def test_missing_file_reports_error_not_exception(self):
        data, err = read_attached_file(os.path.join(self.tmpdir, "does_not_exist.txt"))
        self.assertIsNone(data)
        self.assertIsNotNone(err)


class TestReadAttachedFileImage(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_png_is_base64_encoded_as_image(self):
        path = os.path.join(self.tmpdir, "map_screenshot.png")
        raw = b"\x89PNG\r\n\x1a\nnot a real png but bytes are bytes"
        with open(path, "wb") as f:
            f.write(raw)
        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertTrue(data["is_image"])
        self.assertIsNone(data["text"])
        self.assertEqual(data["mime"], "png")
        self.assertEqual(base64.b64decode(data["b64"]), raw)

    def test_jpg_and_jpeg_both_map_to_jpeg_mime(self):
        for ext in (".jpg", ".jpeg"):
            path = os.path.join(self.tmpdir, f"photo{ext}")
            with open(path, "wb") as f:
                f.write(b"fake jpeg bytes")
            data, err = read_attached_file(path)
            self.assertIsNone(err)
            self.assertEqual(data["mime"], "jpeg")


class TestReadAttachedFileCsv(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_csv_preview_summarizes_shape_and_head(self):
        _skip_if_missing(self, "pandas")
        path = os.path.join(self.tmpdir, "3w.csv")
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write("org,district,sector\nNRC,Aden,Shelter\nUNICEF,Aden,WASH\n")
        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertFalse(data["is_image"])
        self.assertIn("2 rows, 3 columns", data["text"])
        self.assertIn("org", data["text"])

    def test_csv_falls_back_through_encodings_for_cp1252_bytes(self):
        # Excel-exported CSVs are often cp1252 (curly quotes, em dashes), not
        # UTF-8 -- this is the exact scenario the encoding fallback chain
        # exists for.
        _skip_if_missing(self, "pandas")
        path = os.path.join(self.tmpdir, "excel_export.csv")
        with open(path, "wb") as f:
            # U+2019 (curly right single quote) cp1252-encodes to byte 0x92,
            # which is not valid UTF-8 -- this is the exact real-world case
            # (Excel-exported CSVs with "smart quotes") the encoding fallback
            # chain exists for.
            f.write("org,note\nNRC,it’s shelter\n".encode("cp1252"))
        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertIn("1 rows, 2 columns", data["text"])


class TestReadAttachedFileExcel(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_xlsx_preview_summarizes_shape_and_head(self):
        _skip_if_missing(self, "pandas")
        _skip_if_missing(self, "openpyxl")
        import pandas as pd

        path = os.path.join(self.tmpdir, "3w.xlsx")
        pd.DataFrame({"org": ["NRC", "UNICEF"], "district": ["Aden", "Taiz"]}).to_excel(path, index=False)
        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertIn("2 rows, 2 columns", data["text"])


class TestReadAttachedFilePdf(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_extracts_text_from_a_real_pdf(self):
        _skip_if_missing(self, "pypdf")
        _skip_if_missing(self, "reportlab")
        from reportlab.pdfgen import canvas

        path = os.path.join(self.tmpdir, "sitrep.pdf")
        c = canvas.Canvas(path)
        c.drawString(100, 750, "Situation Report: Aden District")
        c.save()

        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertFalse(data["is_image"])
        self.assertIn("Aden", data["text"])


class TestReadAttachedFileDocx(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_extracts_paragraphs_and_table_text_from_a_real_docx(self):
        _skip_if_missing(self, "docx")
        from docx import Document

        path = os.path.join(self.tmpdir, "needs.docx")
        doc = Document()
        doc.add_paragraph("Needs Assessment Summary")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "District"
        table.rows[0].cells[1].text = "Households"
        doc.save(path)

        data, err = read_attached_file(path)
        self.assertIsNone(err)
        self.assertFalse(data["is_image"])
        self.assertIn("Needs Assessment Summary", data["text"])
        self.assertIn("District", data["text"])
        self.assertIn("Households", data["text"])


if __name__ == "__main__":
    unittest.main()
